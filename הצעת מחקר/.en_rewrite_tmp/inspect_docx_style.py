import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def color_value(color):
    if color is None or color.rgb is None:
        return None
    return str(color.rgb)


def length_inches(value):
    return None if value is None else round(value.inches, 4)


def paragraph_spec(paragraph):
    fmt = paragraph.paragraph_format
    first_run = paragraph.runs[0] if paragraph.runs else None
    return {
        "style": paragraph.style.name,
        "text": paragraph.text,
        "alignment": None if paragraph.alignment is None else str(paragraph.alignment),
        "bidi": paragraph._p.pPr is not None
        and paragraph._p.pPr.find(qn("w:bidi")) is not None,
        "space_before_pt": None
        if fmt.space_before is None
        else round(fmt.space_before.pt, 2),
        "space_after_pt": None
        if fmt.space_after is None
        else round(fmt.space_after.pt, 2),
        "line_spacing": fmt.line_spacing,
        "keep_with_next": fmt.keep_with_next,
        "first_run": None
        if first_run is None
        else {
            "font": first_run.font.name,
            "size_pt": None if first_run.font.size is None else first_run.font.size.pt,
            "bold": first_run.bold,
            "italic": first_run.italic,
            "color": color_value(first_run.font.color),
        },
    }


def style_spec(style):
    fmt = style.paragraph_format
    return {
        "font": style.font.name,
        "size_pt": None if style.font.size is None else style.font.size.pt,
        "bold": style.font.bold,
        "italic": style.font.italic,
        "color": color_value(style.font.color),
        "alignment": None if fmt.alignment is None else str(fmt.alignment),
        "space_before_pt": None
        if fmt.space_before is None
        else round(fmt.space_before.pt, 2),
        "space_after_pt": None
        if fmt.space_after is None
        else round(fmt.space_after.pt, 2),
        "line_spacing": fmt.line_spacing,
        "keep_with_next": fmt.keep_with_next,
    }


def table_spec(table):
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    return {
        "rows": len(table.rows),
        "cols": len(table.columns),
        "column_widths_in": [length_inches(col.width) for col in table.columns],
        "cell_text": [[cell.text for cell in row.cells] for row in table.rows],
        "xml_tblW": None if tbl_w is None else tbl_w.get(qn("w:w")),
        "xml_tblW_type": None if tbl_w is None else tbl_w.get(qn("w:type")),
        "grid_dxa": [grid.get(qn("w:w")) for grid in table._tbl.tblGrid.gridCol_lst],
    }


def main():
    path = Path(sys.argv[1])
    doc = Document(path)
    section = doc.sections[0]
    selected_paragraphs = [0, 1, 2, 3, 5, 6, 14, 16, 22, 28, 29, 47, 48]
    data = {
        "path": str(path),
        "sections": [
            {
                "width_in": length_inches(s.page_width),
                "height_in": length_inches(s.page_height),
                "left_in": length_inches(s.left_margin),
                "right_in": length_inches(s.right_margin),
                "top_in": length_inches(s.top_margin),
                "bottom_in": length_inches(s.bottom_margin),
                "header_in": length_inches(s.header_distance),
                "footer_in": length_inches(s.footer_distance),
                "start_type": str(s.start_type),
            }
            for s in doc.sections
        ],
        "styles": {
            name: style_spec(doc.styles[name])
            for name in ["Normal", "Heading 1", "Heading 2"]
        },
        "paragraphs": {
            str(i): paragraph_spec(doc.paragraphs[i]) for i in selected_paragraphs
        },
        "tables": [table_spec(t) for t in doc.tables],
        "header_text": [p.text for p in section.header.paragraphs],
        "footer_text": [p.text for p in section.footer.paragraphs],
    }
    print(json.dumps(data, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
